import io

from openpyxl import load_workbook

from src.constants import (
    PREFERENCE_PREFER_OFF,
    PREFERENCE_PREFER_WORK,
    PREFERENCE_UNAVAILABLE,
    SCHEDULE_STATUS_CONFIRMED,
    SCHEDULE_STATUS_DRAFT,
    SOURCE_MANUAL,
    SOURCE_OPTIMIZED,
)
from src.export_excel import (
    _SHEET_DAILY,
    _SHEET_MONTHLY,
    _SHEET_STAFF_SUMMARY,
    build_schedule_workbook,
    export_schedule_excel,
)
from src.models import (
    DailyRequirementInput,
    MonthlyConditionInput,
    PreferenceInput,
    RoleRequirementInput,
    SchedulerInput,
    StaffInput,
)
from src.month_utils import get_month_dates
from src.repositories import ScheduleAssignmentRecord, ScheduleMonthRecord

YM = "2026-10"  # 31日。10/1 は木曜(3)、10/3 が土曜、10/4 が日曜
DATES = get_month_dates(YM)
LEADER, CHECKER, CLEANER = 1, 2, 3
ROLE_NAMES = {LEADER: "リーダー", CHECKER: "チェッカー", CLEANER: "クリーナー"}


def make_staff(staff_id, role_id=CLEANER, *, active=True, minutes=480, max_consec=5, skill_level=3):
    return StaffInput(
        staff_id=staff_id,
        staff_name=f"スタッフ{staff_id}",
        role_id=role_id,
        daily_work_minutes=minutes,
        max_consecutive_days=max_consec,
        active=active,
        weekday_availability={w: True for w in range(7)},
        skill_level=skill_level,
    )


def make_scheduler_input(staff_list, *, conditions=(), daily_reqs=(), role_reqs=(), prefs=()):
    return SchedulerInput(
        year_month=YM,
        staff=list(staff_list),
        monthly_conditions=list(conditions),
        daily_requirements=list(daily_reqs),
        role_requirements=list(role_reqs),
        preferences=list(prefs),
    )


def assignment(staff_id, work_date, is_working, *, is_locked=False, source=SOURCE_OPTIMIZED):
    return ScheduleAssignmentRecord(
        staff_id=staff_id,
        work_date=work_date,
        is_working=is_working,
        is_locked=is_locked,
        source=source,
    )


def load_bytes(data: bytes):
    return load_workbook(io.BytesIO(data))


# ---------------------------------------------------------------------------
# Sheet structure
# ---------------------------------------------------------------------------


def test_sheet_names_and_order():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    assert wb.sheetnames == [_SHEET_MONTHLY, _SHEET_DAILY, _SHEET_STAFF_SUMMARY]


def test_export_schedule_excel_returns_loadable_bytes():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    data = export_schedule_excel(scheduler_input, [], ROLE_NAMES)
    assert isinstance(data, bytes)
    wb = load_bytes(data)
    assert wb.sheetnames == [_SHEET_MONTHLY, _SHEET_DAILY, _SHEET_STAFF_SUMMARY]


# ---------------------------------------------------------------------------
# T85 月間勤務表
# ---------------------------------------------------------------------------


def test_monthly_sheet_title_and_status_confirmed():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    schedule_month = ScheduleMonthRecord(
        year_month=YM,
        status=SCHEDULE_STATUS_CONFIRMED,
        solver_status="OPTIMAL",
        objective_overstaff=0,
        objective_target_deviation=0,
        objective_prefer_off=0,
        objective_prefer_work=0,
        generated_at="2026-10-01T09:00:00",
        confirmed_at="2026-10-02T10:00:00",
    )
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES, schedule_month)
    ws = wb[_SHEET_MONTHLY]
    assert ws.cell(row=1, column=1).value == "2026年10月 勤務表"
    status_text = ws.cell(row=2, column=1).value
    assert "確定" in status_text
    assert "2026-10-02T10:00:00" in status_text


def test_monthly_sheet_status_draft_and_none():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)

    draft = ScheduleMonthRecord(
        year_month=YM,
        status=SCHEDULE_STATUS_DRAFT,
        solver_status="OPTIMAL",
        objective_overstaff=0,
        objective_target_deviation=0,
        objective_prefer_off=0,
        objective_prefer_work=0,
        generated_at="2026-10-01T09:00:00",
        confirmed_at=None,
    )
    wb_draft = build_schedule_workbook(scheduler_input, [], ROLE_NAMES, draft)
    draft_status = wb_draft[_SHEET_MONTHLY].cell(row=2, column=1).value
    assert "下書き" in draft_status
    assert "確定日時" not in draft_status

    wb_none = build_schedule_workbook(scheduler_input, [], ROLE_NAMES, None)
    none_status = wb_none[_SHEET_MONTHLY].cell(row=2, column=1).value
    assert "未生成" in none_status

    # 3つのラベルはすべて異なる
    assert {draft_status, none_status} != {draft_status}
    confirmed = ScheduleMonthRecord(
        year_month=YM,
        status=SCHEDULE_STATUS_CONFIRMED,
        solver_status="OPTIMAL",
        objective_overstaff=0,
        objective_target_deviation=0,
        objective_prefer_off=0,
        objective_prefer_work=0,
        generated_at="2026-10-01T09:00:00",
        confirmed_at=None,
    )
    wb_confirmed = build_schedule_workbook(scheduler_input, [], ROLE_NAMES, confirmed)
    confirmed_status = wb_confirmed[_SHEET_MONTHLY].cell(row=2, column=1).value
    labels = {draft_status, none_status, confirmed_status}
    assert len(labels) == 3


def test_monthly_sheet_weekday_header_and_shading():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_MONTHLY]

    # C列(3列目) = 10/1 木曜
    assert ws.cell(row=3, column=3).value == 1
    assert ws.cell(row=4, column=3).value == "木"

    # 10/3 が土曜 -> 5列目 (C=1日, D=2日, E=3日)
    sat_col = 3 + (3 - 1)
    sun_col = 3 + (4 - 1)
    assert ws.cell(row=4, column=sat_col).value == "土"
    assert ws.cell(row=4, column=sun_col).value == "日"
    assert ws.cell(row=3, column=sat_col).fill.start_color.rgb == "00DCE6F1"
    assert ws.cell(row=3, column=sun_col).fill.start_color.rgb == "00F2DCDB"

    assert ws.cell(row=3, column=1).value == "スタッフ"
    assert ws.cell(row=3, column=2).value == "ロール"

    last_date_col = 3 + len(DATES) - 1
    assert ws.cell(row=3, column=last_date_col + 1).value == "出勤日数"
    assert ws.cell(row=3, column=last_date_col + 2).value == "勤務時間(h)"


def test_monthly_sheet_working_marks_blank_and_locked_bold():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    day1 = DATES[0]
    day2 = DATES[1]
    assignments = [
        assignment(1, day1, True, is_locked=True),
        assignment(1, day2, False),
        # day3 missing entirely -> treated as off
    ]
    wb = build_schedule_workbook(scheduler_input, assignments, ROLE_NAMES)
    ws = wb[_SHEET_MONTHLY]

    row = 5  # first data row
    col_day1 = 3
    col_day2 = 4
    col_day3 = 5

    cell1 = ws.cell(row=row, column=col_day1)
    assert cell1.value == "○"
    assert cell1.font.bold is True

    cell2 = ws.cell(row=row, column=col_day2)
    assert cell2.value == ""
    assert not cell2.font.bold

    cell3 = ws.cell(row=row, column=col_day3)
    assert cell3.value == ""  # missing assignment == off


def test_monthly_sheet_workdays_and_hours_columns():
    staff_list = [make_staff(1, minutes=480)]
    scheduler_input = make_scheduler_input(staff_list)
    working_dates = DATES[:3]
    assignments = [assignment(1, d, True) for d in working_dates]
    wb = build_schedule_workbook(scheduler_input, assignments, ROLE_NAMES)
    ws = wb[_SHEET_MONTHLY]

    last_date_col = 3 + len(DATES) - 1
    workdays_col = last_date_col + 1
    hours_col = last_date_col + 2

    assert ws.cell(row=5, column=workdays_col).value == 3
    assert ws.cell(row=5, column=hours_col).value == 24.0  # 3 * 480 / 60


def test_monthly_sheet_inactive_staff_excluded_and_freeze_panes():
    staff_list = [make_staff(1, active=True), make_staff(2, active=False)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_MONTHLY]

    names = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert "スタッフ1" in names
    assert "スタッフ2" not in names

    assert ws.freeze_panes == "C5"


def test_monthly_sheet_28_and_31_day_months_column_count():
    ym_28 = "2026-02"  # 2026年は28日
    staff_list = [make_staff(1)]
    scheduler_input_28 = SchedulerInput(year_month=ym_28, staff=staff_list)
    wb_28 = build_schedule_workbook(scheduler_input_28, [], ROLE_NAMES)
    dates_28 = get_month_dates(ym_28)
    assert len(dates_28) == 28
    ws_28 = wb_28[_SHEET_MONTHLY]
    last_date_col_28 = 3 + len(dates_28) - 1
    assert ws_28.cell(row=3, column=last_date_col_28 + 1).value == "出勤日数"

    scheduler_input_31 = make_scheduler_input(staff_list)
    wb_31 = build_schedule_workbook(scheduler_input_31, [], ROLE_NAMES)
    assert len(DATES) == 31
    ws_31 = wb_31[_SHEET_MONTHLY]
    last_date_col_31 = 3 + len(DATES) - 1
    assert ws_31.cell(row=3, column=last_date_col_31 + 1).value == "出勤日数"


# ---------------------------------------------------------------------------
# T86 日別人数
# ---------------------------------------------------------------------------


def test_daily_sheet_headers_and_role_columns():
    staff_list = [make_staff(1, role_id=LEADER), make_staff(2, role_id=CLEANER)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_DAILY]

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers[:6] == ["日付", "曜日", "最低人数", "最大人数", "出勤人数", "過不足"]
    assert "リーダー（必要）" in headers
    assert "リーダー（出勤）" in headers
    assert "チェッカー（必要）" in headers
    assert "クリーナー（出勤）" in headers
    assert headers[-1] == "備考"


def test_daily_sheet_row_values_and_role_counts():
    staff_list = [
        make_staff(1, role_id=LEADER),
        make_staff(2, role_id=CLEANER),
        make_staff(3, role_id=CLEANER),
    ]
    day1 = DATES[0]
    daily_reqs = [
        DailyRequirementInput(work_date=day1, required_total_staff=2, max_total_staff=3, note="繁忙期"),
    ]
    role_reqs = [
        RoleRequirementInput(work_date=day1, role_id=LEADER, required_count=1),
        RoleRequirementInput(work_date=day1, role_id=CLEANER, required_count=1),
    ]
    scheduler_input = make_scheduler_input(
        staff_list, daily_reqs=daily_reqs, role_reqs=role_reqs
    )
    assignments = [
        assignment(1, day1, True),
        assignment(2, day1, True),
        assignment(3, day1, False),
    ]
    wb = build_schedule_workbook(scheduler_input, assignments, ROLE_NAMES)
    ws = wb[_SHEET_DAILY]

    row = 2
    assert ws.cell(row=row, column=1).value == day1
    assert ws.cell(row=row, column=2).value == "木"
    assert ws.cell(row=row, column=3).value == 2  # 最低人数
    assert ws.cell(row=row, column=4).value == 3  # 最大人数
    assert ws.cell(row=row, column=5).value == 2  # 出勤人数
    assert ws.cell(row=row, column=6).value == 0  # 過不足 = 2-2

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    leader_req_col = headers.index("リーダー（必要）") + 1
    leader_att_col = headers.index("リーダー（出勤）") + 1
    cleaner_req_col = headers.index("クリーナー（必要）") + 1
    cleaner_att_col = headers.index("クリーナー（出勤）") + 1
    note_col = headers.index("備考") + 1

    assert ws.cell(row=row, column=leader_req_col).value == 1
    assert ws.cell(row=row, column=leader_att_col).value == 1
    assert ws.cell(row=row, column=cleaner_req_col).value == 1
    assert ws.cell(row=row, column=cleaner_att_col).value == 1
    assert ws.cell(row=row, column=note_col).value == "繁忙期"

    # 2日目は daily_requirements が無い -> 最低人数/最大人数/過不足はブランク
    row2 = 3
    assert ws.cell(row=row2, column=3).value is None
    assert ws.cell(row=row2, column=4).value is None
    assert ws.cell(row=row2, column=6).value is None
    assert ws.cell(row=row2, column=5).value == 0  # 出勤人数は0


def test_daily_sheet_max_total_none_is_blank():
    staff_list = [make_staff(1)]
    day1 = DATES[0]
    daily_reqs = [DailyRequirementInput(work_date=day1, required_total_staff=1, max_total_staff=None)]
    scheduler_input = make_scheduler_input(staff_list, daily_reqs=daily_reqs)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_DAILY]
    assert ws.cell(row=2, column=4).value is None


# ---------------------------------------------------------------------------
# T87 スタッフ集計
# ---------------------------------------------------------------------------


def test_staff_summary_headers():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_STAFF_SUMMARY]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers == [
        "スタッフ",
        "ロール",
        "スキル",
        "1日勤務(分)",
        "出勤日数",
        "勤務時間(分)",
        "所定(分)",
        "最低(分)",
        "最大(分)",
        "目標日数",
        "目標差",
        "前月連勤",
        "希望休(できれば休み)件数",
        "うち出勤",
        "できれば勤務 件数",
        "うち休み",
        "絶対休み件数",
        "うち出勤",
    ]


def test_staff_summary_values_with_condition_and_preferences():
    staff_list = [make_staff(1, minutes=480, skill_level=5)]
    day1, day2, day3, day4 = DATES[0], DATES[1], DATES[2], DATES[3]
    conditions = [
        MonthlyConditionInput(
            staff_id=1,
            year_month=YM,
            target_monthly_minutes=4800,  # 10日相当 (4800/480=10)
            min_monthly_minutes=3000,
            max_monthly_minutes=6000,
            carryover_consecutive_days=2,
        )
    ]
    prefs = [
        PreferenceInput(1, day1, PREFERENCE_PREFER_OFF),  # 希望休だが出勤
        PreferenceInput(1, day2, PREFERENCE_PREFER_WORK),  # できれば勤務だが休み
        PreferenceInput(1, day3, PREFERENCE_UNAVAILABLE),  # 絶対休みだが出勤(違反)
    ]
    scheduler_input = make_scheduler_input(staff_list, conditions=conditions, prefs=prefs)
    assignments = [
        assignment(1, day1, True),
        assignment(1, day2, False),
        assignment(1, day3, True),
        assignment(1, day4, True),
    ]
    wb = build_schedule_workbook(scheduler_input, assignments, ROLE_NAMES)
    ws = wb[_SHEET_STAFF_SUMMARY]

    row = 2
    assert ws.cell(row=row, column=1).value == "スタッフ1"
    assert ws.cell(row=row, column=2).value == "クリーナー"
    assert ws.cell(row=row, column=3).value == 5  # スキル
    assert ws.cell(row=row, column=4).value == 480
    assert ws.cell(row=row, column=5).value == 3  # 出勤日数 (day1, day3, day4)
    assert ws.cell(row=row, column=6).value == 1440  # 3 * 480
    assert ws.cell(row=row, column=7).value == 4800
    assert ws.cell(row=row, column=8).value == 3000
    assert ws.cell(row=row, column=9).value == 6000
    assert ws.cell(row=row, column=10).value == 10  # 目標日数
    assert ws.cell(row=row, column=11).value == -7  # 目標差 = 3 - 10
    assert ws.cell(row=row, column=12).value == 2  # 前月連勤
    assert ws.cell(row=row, column=13).value == 1  # 希望休件数
    assert ws.cell(row=row, column=14).value == 1  # うち出勤
    assert ws.cell(row=row, column=15).value == 1  # できれば勤務件数
    assert ws.cell(row=row, column=16).value == 1  # うち休み
    assert ws.cell(row=row, column=17).value == 1  # 絶対休み件数
    assert ws.cell(row=row, column=18).value == 1  # うち出勤 (違反)


def test_staff_summary_no_condition_is_blank():
    staff_list = [make_staff(1)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_STAFF_SUMMARY]
    row = 2
    for col in (7, 8, 9, 10, 11, 12):
        assert ws.cell(row=row, column=col).value is None
    assert ws.cell(row=row, column=13).value == 0
    assert ws.cell(row=row, column=17).value == 0
    assert ws.cell(row=row, column=18).value == 0


def test_staff_summary_inactive_excluded_and_ordered_by_staff_id():
    staff_list = [make_staff(3), make_staff(1), make_staff(2, active=False)]
    scheduler_input = make_scheduler_input(staff_list)
    wb = build_schedule_workbook(scheduler_input, [], ROLE_NAMES)
    ws = wb[_SHEET_STAFF_SUMMARY]
    names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert names == ["スタッフ1", "スタッフ3"]
