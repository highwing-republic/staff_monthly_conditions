"""シフトのExcel出力（Phase 20, T85〜T87, §50 06_schedule Excel）.

3シートを持つブックを作成する。

- T85 月間勤務表: スタッフ x 日付のグリッド。○=出勤 / 空欄=休み。
- T86 日別人数: 日付ごとの最低人数・出勤人数・ロール別内訳。
- T87 スタッフ集計: スタッフごとの勤務実績・希望反映状況。

DB/Streamlitに依存しない（SchedulerInput・repositoryのレコード型のみを使う）。
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.constants import (
    PREFERENCE_PREFER_OFF,
    PREFERENCE_PREFER_WORK,
    PREFERENCE_UNAVAILABLE,
    SCHEDULE_STATUS_CONFIRMED,
    SCHEDULE_STATUS_DRAFT,
)
from src.models import SchedulerInput
from src.month_utils import get_month_dates, parse_year_month, round_half_up_workdays, weekday_index
from src.repositories import ScheduleAssignmentRecord, ScheduleMonthRecord

_WEEKDAY_LABELS = ("月", "火", "水", "木", "金", "土", "日")

_SHEET_MONTHLY = "月間勤務表"
_SHEET_DAILY = "日別人数"
_SHEET_STAFF_SUMMARY = "スタッフ集計"

_WORKING_MARK = "○"

_HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_SATURDAY_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
_SUNDAY_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
_TITLE_FONT = Font(bold=True, size=14)
_STATUS_FONT = Font(bold=True, size=11)
_HEADER_FONT = Font(bold=True)
_THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_CENTER = Alignment(horizontal="center", vertical="center")


def _active_staff(scheduler_input: SchedulerInput):
    """activeなスタッフをstaff_id昇順で返す."""
    return sorted(
        (s for s in scheduler_input.staff if s.active),
        key=lambda s: s.staff_id,
    )


def _status_label(schedule_month: ScheduleMonthRecord | None) -> str:
    if schedule_month is None:
        return "未生成"
    if schedule_month.status == SCHEDULE_STATUS_CONFIRMED:
        return "確定"
    if schedule_month.status == SCHEDULE_STATUS_DRAFT:
        return "下書き"
    return schedule_month.status


def _year_month_label(year_month: str) -> str:
    year, month = parse_year_month(year_month)
    return f"{year}年{month}月"


def _set_cell(ws: Worksheet, row: int, col: int, value, *, bold: bool = False,
              fill: PatternFill | None = None, align_center: bool = False,
              number_format: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = _THIN_BORDER
    if bold:
        cell.font = Font(bold=True)
    if fill is not None:
        cell.fill = fill
    if align_center:
        cell.alignment = _CENTER
    if number_format is not None:
        cell.number_format = number_format


# ---------------------------------------------------------------------------
# T85 月間勤務表
# ---------------------------------------------------------------------------


def _build_monthly_sheet(
    wb: Workbook,
    scheduler_input: SchedulerInput,
    assignments: list[ScheduleAssignmentRecord],
    role_names: dict[int, str],
    schedule_month: ScheduleMonthRecord | None,
) -> None:
    ws = wb.create_sheet(_SHEET_MONTHLY)

    month_dates = get_month_dates(scheduler_input.year_month)
    staff_list = _active_staff(scheduler_input)
    assignment_by_key: dict[tuple[int, str], ScheduleAssignmentRecord] = {
        (a.staff_id, a.work_date): a for a in assignments
    }

    name_col = 1
    role_col = 2
    first_date_col = 3
    last_date_col = first_date_col + len(month_dates) - 1
    workdays_col = last_date_col + 1
    hours_col = last_date_col + 2

    title_row = 1
    status_row = 2
    header_day_row = 3
    header_weekday_row = 4
    first_data_row = 5

    # --- title / status ---
    status_label = _status_label(schedule_month)
    status_text = f"ステータス: {status_label}"
    if schedule_month is not None and schedule_month.confirmed_at:
        status_text += f"（確定日時: {schedule_month.confirmed_at}）"

    ws.cell(row=title_row, column=1, value=f"{_year_month_label(scheduler_input.year_month)} 勤務表")
    ws.cell(row=title_row, column=1).font = _TITLE_FONT
    ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=hours_col)

    ws.cell(row=status_row, column=1, value=status_text)
    ws.cell(row=status_row, column=1).font = _STATUS_FONT
    ws.merge_cells(start_row=status_row, start_column=1, end_row=status_row, end_column=hours_col)

    # --- header ---
    ws.merge_cells(start_row=header_day_row, start_column=name_col, end_row=header_weekday_row, end_column=name_col)
    _set_cell(ws, header_day_row, name_col, "スタッフ", bold=True, fill=_HEADER_FILL, align_center=True)

    ws.merge_cells(start_row=header_day_row, start_column=role_col, end_row=header_weekday_row, end_column=role_col)
    _set_cell(ws, header_day_row, role_col, "ロール", bold=True, fill=_HEADER_FILL, align_center=True)

    for offset, work_date in enumerate(month_dates):
        col = first_date_col + offset
        wd = weekday_index(work_date)
        day_number = int(work_date[8:10])
        fill = _SATURDAY_FILL if wd == 5 else _SUNDAY_FILL if wd == 6 else _HEADER_FILL
        _set_cell(ws, header_day_row, col, day_number, bold=True, fill=fill, align_center=True)
        _set_cell(ws, header_weekday_row, col, _WEEKDAY_LABELS[wd], bold=True, fill=fill, align_center=True)

    ws.merge_cells(start_row=header_day_row, start_column=workdays_col, end_row=header_weekday_row, end_column=workdays_col)
    _set_cell(ws, header_day_row, workdays_col, "出勤日数", bold=True, fill=_HEADER_FILL, align_center=True)

    ws.merge_cells(start_row=header_day_row, start_column=hours_col, end_row=header_weekday_row, end_column=hours_col)
    _set_cell(ws, header_day_row, hours_col, "勤務時間(h)", bold=True, fill=_HEADER_FILL, align_center=True)

    # --- staff rows ---
    for row_offset, staff in enumerate(staff_list):
        row = first_data_row + row_offset
        _set_cell(ws, row, name_col, staff.staff_name)
        _set_cell(ws, row, role_col, role_names.get(staff.role_id, str(staff.role_id)))

        workday_count = 0
        for offset, work_date in enumerate(month_dates):
            col = first_date_col + offset
            record = assignment_by_key.get((staff.staff_id, work_date))
            is_working = record is not None and record.is_working
            is_locked = record is not None and record.is_locked
            value = _WORKING_MARK if is_working else ""
            if is_working:
                workday_count += 1
            _set_cell(ws, row, col, value, bold=is_locked, align_center=True)

        _set_cell(ws, row, workdays_col, workday_count, align_center=True)
        hours = round(workday_count * staff.daily_work_minutes / 60, 1)
        _set_cell(ws, row, hours_col, hours, align_center=True, number_format="0.0")

    # --- layout ---
    ws.freeze_panes = ws.cell(row=first_data_row, column=first_date_col).coordinate
    ws.column_dimensions[get_column_letter(name_col)].width = 14
    ws.column_dimensions[get_column_letter(role_col)].width = 10
    for offset in range(len(month_dates)):
        ws.column_dimensions[get_column_letter(first_date_col + offset)].width = 4
    ws.column_dimensions[get_column_letter(workdays_col)].width = 10
    ws.column_dimensions[get_column_letter(hours_col)].width = 12


# ---------------------------------------------------------------------------
# T86 日別人数
# ---------------------------------------------------------------------------


def _build_daily_sheet(
    wb: Workbook,
    scheduler_input: SchedulerInput,
    assignments: list[ScheduleAssignmentRecord],
    role_names: dict[int, str],
) -> None:
    ws = wb.create_sheet(_SHEET_DAILY)

    month_dates = get_month_dates(scheduler_input.year_month)
    staff_by_id = {s.staff_id: s for s in _active_staff(scheduler_input)}
    active_ids = set(staff_by_id.keys())

    working_by_date: dict[str, set[int]] = {d: set() for d in month_dates}
    for record in assignments:
        if record.is_working and record.staff_id in active_ids and record.work_date in working_by_date:
            working_by_date[record.work_date].add(record.staff_id)

    daily_req_by_date = {r.work_date: r for r in scheduler_input.daily_requirements}
    role_req_by_key = {
        (r.work_date, r.role_id): r.required_count for r in scheduler_input.role_requirements
    }
    role_ids = sorted(role_names.keys())

    headers = ["日付", "曜日", "最低人数", "最大人数", "出勤人数", "過不足"]
    for role_id in role_ids:
        role_name = role_names[role_id]
        headers.append(f"{role_name}（必要）")
        headers.append(f"{role_name}（出勤）")
    headers.append("備考")

    for col, header in enumerate(headers, start=1):
        _set_cell(ws, 1, col, header, bold=True, fill=_HEADER_FILL, align_center=True)

    for row_offset, work_date in enumerate(month_dates):
        row = 2 + row_offset
        wd = weekday_index(work_date)
        req = daily_req_by_date.get(work_date)
        required_total = req.required_total_staff if req is not None else None
        max_total = req.max_total_staff if req is not None else None
        note = req.note if req is not None else None

        working_ids = working_by_date[work_date]
        attended_total = len(working_ids)
        shortfall = attended_total - required_total if required_total is not None else None

        fill = _SATURDAY_FILL if wd == 5 else _SUNDAY_FILL if wd == 6 else None

        col = 1
        _set_cell(ws, row, col, work_date, fill=fill, align_center=True)
        col += 1
        _set_cell(ws, row, col, _WEEKDAY_LABELS[wd], fill=fill, align_center=True)
        col += 1
        _set_cell(ws, row, col, required_total, fill=fill, align_center=True)
        col += 1
        _set_cell(ws, row, col, max_total, fill=fill, align_center=True)
        col += 1
        _set_cell(ws, row, col, attended_total, fill=fill, align_center=True)
        col += 1
        _set_cell(ws, row, col, shortfall, fill=fill, align_center=True)
        col += 1

        for role_id in role_ids:
            # §27: ロール別必要人数の行がなければ0
            required_role = role_req_by_key.get((work_date, role_id), 0)
            attended_role = sum(
                1 for staff_id in working_ids if staff_by_id[staff_id].role_id == role_id
            )
            _set_cell(ws, row, col, required_role, fill=fill, align_center=True)
            col += 1
            _set_cell(ws, row, col, attended_role, fill=fill, align_center=True)
            col += 1

        _set_cell(ws, row, col, note, fill=fill)

    ws.freeze_panes = "A2"
    ws.column_dimensions[get_column_letter(1)].width = 12
    ws.column_dimensions[get_column_letter(2)].width = 6
    for i in range(3, len(headers)):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.column_dimensions[get_column_letter(len(headers))].width = 20


# ---------------------------------------------------------------------------
# T87 スタッフ集計
# ---------------------------------------------------------------------------


def _build_staff_summary_sheet(
    wb: Workbook,
    scheduler_input: SchedulerInput,
    assignments: list[ScheduleAssignmentRecord],
    role_names: dict[int, str],
) -> None:
    ws = wb.create_sheet(_SHEET_STAFF_SUMMARY)

    month_dates = set(get_month_dates(scheduler_input.year_month))
    staff_list = _active_staff(scheduler_input)
    condition_by_staff = {c.staff_id: c for c in scheduler_input.monthly_conditions}

    working_dates_by_staff: dict[int, set[str]] = {s.staff_id: set() for s in staff_list}
    for record in assignments:
        if (
            record.is_working
            and record.staff_id in working_dates_by_staff
            and record.work_date in month_dates
        ):
            working_dates_by_staff[record.staff_id].add(record.work_date)

    prefs_by_staff: dict[int, list] = {s.staff_id: [] for s in staff_list}
    for pref in scheduler_input.preferences:
        if pref.staff_id in prefs_by_staff and pref.work_date in month_dates:
            prefs_by_staff[pref.staff_id].append(pref)

    headers = [
        "スタッフ",
        "ロール",
        "スキル",
        "1日勤務(分)",
        "出勤日数",
        "勤務時間(分)",
        "所定(分)",
        "最低(分)",
        "最大(分)",
        "目標日数",
        "目標差",
        "前月連勤",
        "希望休(できれば休み)件数",
        "うち出勤",
        "できれば勤務 件数",
        "うち休み",
        "絶対休み件数",
        "うち出勤",
    ]
    for col, header in enumerate(headers, start=1):
        _set_cell(ws, 1, col, header, bold=True, fill=_HEADER_FILL, align_center=True)

    for row_offset, staff in enumerate(staff_list):
        row = 2 + row_offset
        working_dates = working_dates_by_staff[staff.staff_id]
        workday_count = len(working_dates)
        minutes_worked = workday_count * staff.daily_work_minutes

        cond = condition_by_staff.get(staff.staff_id)
        target_minutes = cond.target_monthly_minutes if cond is not None else None
        min_minutes = cond.min_monthly_minutes if cond is not None else None
        max_minutes = cond.max_monthly_minutes if cond is not None else None
        carryover = cond.carryover_consecutive_days if cond is not None else None

        if target_minutes is not None:
            target_workdays = round_half_up_workdays(target_minutes, staff.daily_work_minutes)
            target_diff = workday_count - target_workdays
        else:
            target_workdays = None
            target_diff = None

        staff_prefs = prefs_by_staff[staff.staff_id]
        prefer_off = [p for p in staff_prefs if p.preference_type == PREFERENCE_PREFER_OFF]
        prefer_off_worked = sum(1 for p in prefer_off if p.work_date in working_dates)
        prefer_work = [p for p in staff_prefs if p.preference_type == PREFERENCE_PREFER_WORK]
        prefer_work_off = sum(1 for p in prefer_work if p.work_date not in working_dates)
        unavailable = [p for p in staff_prefs if p.preference_type == PREFERENCE_UNAVAILABLE]
        unavailable_worked = sum(1 for p in unavailable if p.work_date in working_dates)

        values = [
            staff.staff_name,
            role_names.get(staff.role_id, str(staff.role_id)),
            staff.skill_level,
            staff.daily_work_minutes,
            workday_count,
            minutes_worked,
            target_minutes,
            min_minutes,
            max_minutes,
            target_workdays,
            target_diff,
            carryover,
            len(prefer_off),
            prefer_off_worked,
            len(prefer_work),
            prefer_work_off,
            len(unavailable),
            unavailable_worked,
        ]
        for col, value in enumerate(values, start=1):
            align = col > 2
            _set_cell(ws, row, col, value, align_center=align)

    ws.freeze_panes = "A2"
    ws.column_dimensions[get_column_letter(1)].width = 14
    ws.column_dimensions[get_column_letter(2)].width = 10
    for i in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def build_schedule_workbook(
    scheduler_input: SchedulerInput,
    assignments: list[ScheduleAssignmentRecord],
    role_names: dict[int, str],
    schedule_month: ScheduleMonthRecord | None = None,
) -> Workbook:
    """T85〜T87の3シートを持つExcelワークブックを作成する."""
    wb = Workbook()
    # デフォルトの空シートを削除してから、順序通りに作成する。
    default_sheet = wb.active
    wb.remove(default_sheet)

    _build_monthly_sheet(wb, scheduler_input, assignments, role_names, schedule_month)
    _build_daily_sheet(wb, scheduler_input, assignments, role_names)
    _build_staff_summary_sheet(wb, scheduler_input, assignments, role_names)

    return wb


def export_schedule_excel(
    scheduler_input: SchedulerInput,
    assignments: list[ScheduleAssignmentRecord],
    role_names: dict[int, str],
    schedule_month: ScheduleMonthRecord | None = None,
) -> bytes:
    """st.download_button向けにワークブックをbytesとして返す."""
    wb = build_schedule_workbook(scheduler_input, assignments, role_names, schedule_month)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
